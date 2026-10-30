"""Journal-style formatting of the pandoc DOCX (fonts, spacing, table borders), then PDF export."""
import subprocess
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

d = Document('paper/manuscript.docx')
for s in d.sections:
    s.left_margin = s.right_margin = Cm(2.2); s.top_margin = s.bottom_margin = Cm(2.0)
def font(st, size, bold=None, color=None):
    st.font.name = 'Times New Roman'; st.font.size = Pt(size)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    if bold is not None: st.font.bold = bold
    if color:
        st.font.color.rgb = RGBColor(*color)
        for c in st.element.rPr.findall(qn('w:color')):
            for att in ('w:themeColor', 'w:themeShade', 'w:themeTint'):
                if c.get(qn(att)) is not None: del c.attrib[qn(att)]
for name, size, bold in [('Normal', 11, None), ('Body Text', 11, None), ('First Paragraph', 11, None), ('Title', 17, True),
                         ('Heading 1', 13.5, True), ('Heading 2', 12, True), ('Heading 3', 11, True), ('Compact', 9, None),
                         ('Image Caption', 9.5, None), ('Table Caption', 9.5, None), ('Author', 11, None)]:
    try: font(d.styles[name], size, bold, (0, 0, 0) if name.startswith('Heading') or name == 'Title' else None)
    except KeyError: pass
for name in ('Body Text', 'First Paragraph'):
    try:
        pf = d.styles[name].paragraph_format; pf.line_spacing = 1.15; pf.space_after = Pt(6)
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    except KeyError: pass
def border(tbl):
    tblPr = tbl._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for e, sz in (('top', 12), ('bottom', 12), ('insideH', 4)):
        x = OxmlElement(f'w:{e}'); x.set(qn('w:val'), 'single'); x.set(qn('w:sz'), str(sz)); x.set(qn('w:color'), '000000'); b.append(x)
    tblPr.append(b)
for t in d.tables:
    border(t)
    for i, row in enumerate(t.rows):
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(1)
                for r in p.runs:
                    r.font.size = Pt(8.5); r.font.name = 'Times New Roman'
                    if i == 0: r.font.bold = True
for p in d.paragraphs:
    if p.style.name.startswith('Heading') or p.style.name == 'Title':
        for r in p.runs: r.font.color.rgb = RGBColor(0x1F, 0x2A, 0x44)
d.save('paper/manuscript.docx')
subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', 'paper', 'paper/manuscript.docx'],
               check=True, capture_output=True)
print('ok')
